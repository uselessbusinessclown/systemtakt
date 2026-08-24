
import os, sys
sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'lib'))
sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))), 'src', 'lib'))
from paths import ROOT, DATA, OUT, DB, GTFS  # noqa: E402
import sqlite3, pandas as pd, os
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

con=sqlite3.connect(DB)
OUT=os.path.join(ROOT, 'Deutschland_Fernverkehr_Datenbank_und_Systemtakt.xlsx')

SHEETS=[
 ('Bahnhöfe', """SELECT b.name AS Bahnhof, b.stadt AS Stadt, b.land AS Land, b.bundesland AS Bundesland,
    b.lat AS Breite, b.lon AS Laenge, b.halte_tag AS "Halte/Tag gesamt", b.halte_ice AS "davon ICE",
    b.halte_sprinter AS "davon ICE-Sprinter", b.halte_ic AS "davon IC/EC",
    b.linien AS "Linien", b.rang AS "Rang", b.linien_liste AS "Linien im Detail"
    FROM bahnhof_bedienung b ORDER BY b.halte_tag DESC"""),
 ('Linien ICE', """SELECT linie_key AS Linie, nummer AS Nr, betreiber AS Betreiber, bezeichnung AS Relation,
    fahrten_tag AS "Fahrten/Tag", fahrten_tag_sprinter AS "davon Sprinter", halte_typisch AS "Halte (typ.)",
    fahrzeit_min AS "Fahrzeit min", km_luftlinie AS "km (Luftlinie)", vschnitt AS "Ø km/h", laufweg AS Laufweg
    FROM linie WHERE produkt='ICE' ORDER BY fahrten_tag DESC"""),
 ('Linien IC-EC', """SELECT linie_key AS Linie, nummer AS Nr, betreiber AS Betreiber, bezeichnung AS Relation,
    fahrten_tag AS "Fahrten/Tag", halte_typisch AS "Halte (typ.)", fahrzeit_min AS "Fahrzeit min",
    km_luftlinie AS "km (Luftlinie)", vschnitt AS "Ø km/h", laufweg AS Laufweg
    FROM linie WHERE produkt='IC' ORDER BY fahrten_tag DESC"""),
 ('Zugfahrten', """SELECT trip_id AS ID, gattung AS Gattung, linie AS Linie, kategorie AS Kategorie,
    sprinter_korridor AS "Sprinter-Korridor", betreiber AS Betreiber, von AS Von, nach AS Nach,
    von_stadt AS "Von (Stadt)", nach_stadt AS "Nach (Stadt)",
    printf('%02d:%02d',abfahrt/3600%24,abfahrt%3600/60) AS Abfahrt,
    printf('%02d:%02d',ankunft/3600%24,ankunft%3600/60) AS Ankunft,
    dauer_min AS "Dauer min", halte AS Halte, km_luftlinie AS "km (Luftlinie)", vschnitt_kmh AS "Ø km/h",
    verkehrstage AS "Verkehrstage (31 T.)", inland AS "nur Inland"
    FROM trip_summary ORDER BY gattung, linie, abfahrt"""),
 ('Kanten Ist', """SELECT a AS "Von", b AS "Nach", km AS "km (Luftlinie)", t_min AS "Fahrzeit min (Bestwert)",
    t_median AS "Median min", t_max AS "Max min", fahrten_tag AS "Fahrten/Tag", produkte AS Produkte
    FROM kante ORDER BY fahrten_tag DESC"""),
 ('ITF Knotenzeiten', """SELECT name AS Vollknoten, knotenzeit AS Knotenzeit, ankunft AS Ankunftsfenster,
    abfahrt AS Abfahrtsfenster, grad AS "Kanten (Grad)", gewicht AS "Gewicht (Prio×Fahrten)",
    off_a1 AS "Ankunft Ri1 (min vor Knoten)", off_d1 AS "Abfahrt Ri1 (min nach)",
    off_a2 AS "Ankunft Ri2", off_d2 AS "Abfahrt Ri2" FROM itf_knotenzeit ORDER BY gewicht DESC"""),
 ('ITF Kanten Soll', """SELECT a AS "Knoten A", b AS "Knoten B", t_basis AS "Systemfahrzeit min",
    t_soll_r1 AS "Soll A→B min", t_soll_r2 AS "Soll B→A min", zuschlag_r1 AS "Zuschlag A→B",
    zuschlag_r2 AS "Zuschlag B→A", reserve_r1 AS "Reserve A→B %", reserve_r2 AS "Reserve B→A %",
    massnahme AS Bewertung, gewicht AS Gewicht, f_sprinter AS "Sprinter/Tag", f_ice AS "ICE/Tag", f_ic AS "IC/Tag"
    FROM itf_ergebnis ORDER BY gewicht DESC"""),
 ('Knotenzahl', """SELECT variante AS Variante, vollknoten AS Vollknoten, kanten AS Netzkanten,
    reserve_pct AS "Ø Reserve %", sollband_pct AS "Sollband %",
    reisezeit_verlaengerung_min AS "Reisezeit +min", bemerkung AS Bemerkung FROM knotenzahl ORDER BY vollknoten"""),
 ('Fahrplan Züge', """SELECT zug_id AS "Zug-Nr.", linie AS Linie, kategorie AS Produktebene, richtung AS Richtung,
    von AS Von, nach AS Nach, printf('%02d:%02d',abfahrt/60,abfahrt%60) AS Abfahrt,
    printf('%02d:%02d',ankunft/60,ankunft%60) AS Ankunft, dauer AS "Dauer min", halte AS Halte, takt AS "Takt min"
    FROM fahrplan_zug ORDER BY zug_id"""),
 ('Fahrplan Halte', """SELECT h.zug_id AS "Zug-Nr.", z.linie AS Linie, h.folge AS Folge, h.station AS Bahnhof,
    CASE WHEN h.an IS NULL THEN '' ELSE printf('%02d:%02d',h.an/60,h.an%60) END AS Ankunft,
    CASE WHEN h.ab IS NULL THEN '' ELSE printf('%02d:%02d',h.ab/60,h.ab%60) END AS Abfahrt,
    CASE h.knoten WHEN 1 THEN 'Vollknoten' ELSE '' END AS Rolle
    FROM fahrplan_halt h JOIN fahrplan_zug z ON z.zug_id=h.zug_id ORDER BY h.zug_id,h.folge"""),
 ('Zielliniennetz', """SELECT prio AS Prio, ebene AS Ebene, linie_key AS Linie, bezeichnung AS Relation,
    betreiber AS Betreiber, richtungsfahrten_ist AS "Ist Fahrten/Ri/Tag", takt_ist AS "Takt heute",
    takt_ziel AS "Takt Szenario A", richtungsfahrten_ziel AS "A Fahrten/Ri/Tag",
    takt_ziel_b AS "Takt Szenario B", richtungsfahrten_ziel_b AS "B Fahrten/Ri/Tag",
    km AS "Strecken-km", knotenfolge AS "Knotenfolge mit Knotenminute"
    FROM ziel_linie ORDER BY prio, richtungsfahrten_ist DESC"""),
 ('Sprinter-Korridore', """SELECT korridor AS Korridor, relation AS Relation, km AS "Strecken-km",
    fahrten_tag_ist AS "Fahrten/Tag heute", takt_ziel AS "Takt Ziel", richtungsfahrten_ziel AS "Ziel Fahrten/Ri/Tag",
    fahrzeit_ist_min AS "Fahrzeit heute min", fahrzeit_ziel_min AS "Fahrzeit Ziel min", knotenlage AS Knotenlage
    FROM ziel_sprinter"""),
 ('Knotenkapazität', """SELECT knoten AS Knoten, knotenzeit AS Knotenzeit, grad AS "Kanten",
    zuege_tag AS "Halte/Tag", zuege_hvz_h AS "Züge/h HVZ", zuege_je_knotenfenster AS "Züge je Knotenfenster",
    gleise_bedarf AS "Bahnsteiggleise (Bedarf FV)", bewertung AS Bewertung
    FROM knoten_kapazitaet ORDER BY zuege_je_knotenfenster DESC"""),
 ('Takt heute', """SELECT linie AS Linie, knoten AS Knoten, richtung AS "Richtung (Ziel)",
    abfahrten_tag AS "Abfahrten/Tag", haeufigste_minute AS "häufigste Minute",
    anteil_haeufigste AS "Anteil", top2_anteil AS "Anteil Top-2", streuung AS "Streuung min",
    takt_bewertung AS Bewertung, minuten AS "Minutenverteilung"
    FROM ist_takt ORDER BY abfahrten_tag DESC"""),
 ('Taktvarianten', """SELECT variante AS Variante, beschreibung AS Beschreibung, verfahren AS Lösungsverfahren,
    knotenraster AS Knotenraster, beschleunigung_erlaubt AS "Beschleunigung erlaubt",
    kanten_mit_beschleunigung AS "Kanten mit Verkürzung", kosten AS Zielfunktion,
    kosten_je_gewicht AS "Zielfunktion je Gewicht", sollband_pct AS "Anteil im Sollband %",
    mittlere_reserve_pct AS "Ø Fahrzeitreserve %", max_zuschlag_min AS "größter Zuschlag min",
    bewertung AS Bewertung FROM takt_varianten ORDER BY variante"""),
 ('Zentralität', """SELECT bahnhof AS Bahnhof, grad AS "Netzkanten", gewicht AS Gewicht, halte_tag AS "Halte/Tag",
    zwischenzentralitaet AS "Zwischenzentralität", rang_zentralitaet AS "Rang Zentralität",
    mittlere_fahrzeit_zum_netz AS "Ø Fahrzeit zum Netz (min)", rang_erreichbarkeit AS "Rang Erreichbarkeit"
    FROM zentralitaet ORDER BY zwischenzentralitaet DESC"""),
 ('Zentrierung', """SELECT bahnhof AS "Prioritätsknoten", kanten AS Netzkanten, gewicht AS Gewicht,
    reserve_basis AS "Reserve ohne Zentrierung %", reserve_zentriert AS "Reserve mit Zentrierung %",
    gewinn_knoten AS "Gewinn am Knoten (Prozentpunkte)", sollband_basis AS "Sollband ohne %",
    sollband_zentriert AS "Sollband mit %", netz_reserve_basis AS "Netz-Reserve ohne %",
    netz_reserve_zentriert AS "Netz-Reserve mit %", netz_kosten_effekt AS "Effekt aufs Netz (Prozentpunkte)",
    optimal AS "Optimalität bewiesen", bewertung AS Bewertung FROM zentrierung ORDER BY gewinn_knoten DESC"""),
 ('Frankfurt Varianten', """SELECT variante AS Variante, beschreibung AS Beschreibung, kanten AS "ITF-Kanten",
    gewicht AS "Gewicht", kosten AS "Zielfunktion", kosten_je_gewicht AS "Kosten je Gewichtseinheit",
    beschl_kanten AS "Kanten mit Beschleunigungsbedarf", sollband_pct AS "Anteil im Sollband %",
    mittlere_reserve_pct AS "Ø Reserve %", hbf_knotenminute AS "Knotenminute Hbf",
    flughafen_knotenminute AS "Knotenminute Flughafen", hbf_gleisminuten AS "Gleisminuten/Tag Hbf",
    flughafen_gleisminuten AS "Gleisminuten/Tag Flughafen", flughafen_auslastung_pct AS "Auslastung Flughafen (4 Gleise) %",
    bewertung AS Bewertung FROM ffm_varianten"""),
 ('Frankfurt Belegung', """SELECT bahnhof AS Bahnhof, zugart AS Zugart, kategorie AS Produktklasse,
    fahrten_tag AS "Fahrten/Tag", gleisminuten_tag AS "Gleisminuten/Tag", anteil_pct AS "Anteil %"
    FROM ffm_belegung"""),
 ('Frankfurt Linien', """SELECT linie AS Linie, fahrten_tag AS "Fahrten/Tag",
    CASE haelt_hbf WHEN 1 THEN 'ja' ELSE '–' END AS "hält Hbf",
    CASE haelt_flughafen WHEN 1 THEN 'ja' ELSE '–' END AS "hält Flughafen",
    verlust AS "Erreichbarkeit" FROM ffm_linien ORDER BY fahrten_tag DESC"""),
 ('Betriebsleistung', """SELECT ebene AS Ebene, zugkm_tag_ist AS "Zug-km/Tag heute", zugkm_tag_A AS "Szenario A",
    zugkm_tag_B AS "Szenario B", mio_jahr_ist AS "Mio Zug-km/Jahr heute", mio_jahr_A AS "A Mio/Jahr",
    mio_jahr_B AS "B Mio/Jahr", delta_A_pct AS "Δ A %", delta_B_pct AS "Δ B %"
    FROM betriebsleistung"""),
]
info=pd.DataFrame({
 'Feld':['Titel','Datenstand','Datenquelle','Abdeckung','Zugfahrten/Tag','Bahnhöfe gesamt','davon in Deutschland',
         'Linien','ICE-Sprinter','Methodik Sprinter','Methodik Systemtakt','Symmetrieminute','Knotenaufenthalt',
         'Szenario A','Szenario B','Frankfurt-Varianten','Neubaufreier Takt','Lösungsverfahren','Zentrierung','Zahl der Vollknoten','Fahrplan','Einschränkung 1','Einschränkung 2','Erstellt'],
 'Wert':['Datenbank des deutschen Schienenfernverkehrs (ICE/IC/EC) und Systemtakt-Planung',
         'Fahrplan 22.08.2026 – 21.09.2026',
         'DELFI e.V. / gtfs.de, GTFS-Feed „Fernverkehr Deutschland“ (de_fv), abgerufen 22.08.2026',
         'alle Fern­verkehrszüge mit Halt in Deutschland, inkl. ausländischer Abschnitte',
         'rund 1.090 (Mo–Fr)','567','290','89 (ICE/ECE, IC/EC, RJ/EN)',
         'rund 72 Fahrten/Tag auf 7 offiziellen Korridoren',
         'Der offene Feed enthält keine Zugnummern. Sprinter werden über die sieben von der DB vermarkteten '
         'Korridore identifiziert und dort über das Fahrzeit-Ranking auf die von bahn.de genannte Angebotsdichte kalibriert.',
         'Phasenoptimierung (PESP-Ansatz): jedem Knoten wird eine Knotenminute im 30-Minuten-Raster zugewiesen; '
         'Zielfunktion minimiert gewichtete Abweichung der Kantenfahrzeiten. Gewichte 3× Sprinter, 2× ICE, 1× IC.',
         'Minute 0 (Ankünfte :57/:27, Abfahrten :03/:33), Anker Berlin Hbf = :00',
         '6 Minuten (Ankunft bis Abfahrt) als Übergangszeit im Knoten',
         'Taktordnung – vorhandenes Angebot auf die nächstgelegene saubere Taktstufe gebracht',
         'Vollausbau ITF – mindestens 120-min-Takt, Hauptachsen 30 min',
         'Geprüft wurde, ob Fernverkehr vom Hauptbahnhof an den Flughafen-Fernbahnhof verlagert werden sollte. '
         'Verglichen werden Status quo, Variante A (alles an den Flughafen), Variante B (Sprinter an den Flughafen) '
         'und Variante C (Fernbahntunnel). Maßstäbe: Taktpassung, Gleisbelegung, Erreichbarkeit.',
         'Der ausgewiesene Knotenplan kommt ohne jede Fahrzeitverkürzung aus: keine Kante muss schneller '
         'werden als sie heute schon gefahren wird. Der Takt entsteht allein aus Fahrzeitzuschlägen.',
         'Der Knotenplan ist nicht mehr heuristisch, sondern als gemischt-ganzzahliges Programm exakt gelöst '
         '(HiGHS über scipy.optimize.milp). Die Optimalität ist bewiesen.',
         'Geprüft wurde, ob eine Zentrierung des Takts auf einen besonders zentralen Großbahnhof hilft. '
         'Die Wahl des Ankerbahnhofs ist wirkungslos (reine Drehung). Wirksam ist dagegen die Höhergewichtung '
         'eines Prioritätsknotens in der Optimierung.',
         'Entscheidend für die Reisezeit ist, wie viele Bahnhöfe Vollknoten sind. Jeder zusätzliche Vollknoten '
         'kostet rund eine halbe Minute mittlere Reisezeit. Gewählt wurden 20 Vollknoten; alle übrigen Bahnhöfe '
         'sind Durchgangshalte ohne Knotenbindung.',
         'Aus dem Knotenplan ist ein vollständiger Tagesfahrplan abgeleitet: 1.700 Züge, 12.876 Halte, '
         'Betriebszeit 5–23 Uhr. Er liegt als Kursbuch der Linienfahrpläne und als Band der Bahnhofsfahrpläne bei.',
         'Fahrzeiten sind Ist-Fahrzeiten aus dem Fahrplan, keine Simulationswerte; Streckenkilometer sind '
         'Luftlinie × 1,20 (Näherung).',
         'Kapazitätsaussagen sind Grobabschätzungen aus Halten je Knotenfenster, keine Betriebssimulation.',
         '23.08.2026 · Fassung 4: 20 Vollknoten, beidseitige Knotenbindung, vollständiger Fahrplan']})

with pd.ExcelWriter(OUT, engine='openpyxl') as xw:
    info.to_excel(xw, sheet_name='Lies mich', index=False)
    for name,q in SHEETS:
        pd.read_sql_query(q, con).to_excel(xw, sheet_name=name, index=False)

wb=load_workbook(OUT)
HEAD=PatternFill('solid', fgColor='1F3864'); HF=Font(name='Arial', size=10, bold=True, color='FFFFFF')
BF=Font(name='Arial', size=10)
thin=Side(style='thin', color='D9D9D9')
for ws in wb.worksheets:
    ws.freeze_panes='A2'
    maxc=ws.max_column
    for c in range(1,maxc+1):
        cell=ws.cell(1,c); cell.fill=HEAD; cell.font=HF
        cell.alignment=Alignment(vertical='center', wrap_text=True)
        w=max([len(str(ws.cell(r,c).value or '')) for r in range(1,min(ws.max_row,400)+1)])
        ws.column_dimensions[get_column_letter(c)].width=min(max(11,w+2),52)
    ws.row_dimensions[1].height=32
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font=BF; cell.border=Border(bottom=thin)
            if isinstance(cell.value,str) and len(cell.value)>60:
                cell.alignment=Alignment(wrap_text=False)
    if ws.max_row>1:
        ws.auto_filter.ref=f"A1:{get_column_letter(maxc)}{ws.max_row}"
ws=wb['Lies mich']; ws.column_dimensions['A'].width=26; ws.column_dimensions['B'].width=110
for r in range(2, ws.max_row+1): ws.cell(r,2).alignment=Alignment(wrap_text=True, vertical='top')
wb.save(OUT)
print('geschrieben:', OUT, os.path.getsize(OUT)//1024, 'KB')
print('Blätter:', wb.sheetnames)
